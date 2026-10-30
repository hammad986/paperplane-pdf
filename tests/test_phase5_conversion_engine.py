"""
Paperplane PDF — Phase 5 Automated Test Suite
================================================
Comprehensive tests for:
1. Conversion Capability Matrix (dynamic engine detection, status reporting, API)
2. Image Conversions (JPG, PNG, WebP, multi-image batch, page sizes, fit modes, PDF to images ZIP)
3. Text Conversions (TXT to PDF, PDF to TXT)
4. Markdown Conversions (Markdown to PDF with GFM tables/CSS, PDF to Markdown structure)
5. HTML Conversions (HTML to PDF, PDF to HTML sections, SSRF security guards)
6. Office Conversions (DOCX to PDF, XLSX to PDF, PPTX to PDF, PDF to DOCX, PDF to XLSX tables, PDF to PPTX slides)
7. PDF/A Archival Conversion
8. Validation & Security (file signatures, container integrity, temp cleanup, process queue)
"""

import io
import json
import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx
import fitz  # PyMuPDF
import openpyxl
import pptx
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw

import converters
import server


class TestPhase5ConversionEngine(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(server.app)
        cls.temp_dir = Path(tempfile.mkdtemp())

        # 1. Create a rich multi-page test PDF with headings, paragraphs, and tables
        cls.sample_pdf = cls.temp_dir / "sample_report.pdf"
        doc = fitz.open()

        # Page 1: Headings and paragraphs
        p1 = doc.new_page(width=595.28, height=841.89)
        p1.insert_text((72, 80), "Paperplane Executive Summary 2026", fontsize=22)
        p1.insert_text((72, 120), "Section 1: Architecture Overview", fontsize=16)
        p1.insert_text((72, 160), "Paperplane is built with local-first processing and zero mandatory uploads.", fontsize=11)
        p1.insert_text((72, 185), "- Multi-engine document conversion pipeline", fontsize=11)
        p1.insert_text((72, 205), "- High-performance vector and raster processing", fontsize=11)

        # Page 2: Table
        p2 = doc.new_page(width=595.28, height=841.89)
        p2.insert_text((72, 80), "Financial Performance", fontsize=18)
        # Draw a table with lines so find_tables() can detect it
        table_x, table_y = 72, 120
        row_h = 24
        cols = [100, 150, 150]
        rows_data = [
            ["Metric", "Target", "Achieved"],
            ["Throughput", "1,000 docs/sec", "2,500 docs/sec"],
            ["Uptime", "99.9%", "99.99%"]
        ]
        # Draw grid
        tot_w = sum(cols)
        tot_h = row_h * len(rows_data)
        for r_idx in range(len(rows_data) + 1):
            curr_y = table_y + r_idx * row_h
            p2.draw_line(fitz.Point(table_x, curr_y), fitz.Point(table_x + tot_w, curr_y))
        curr_x = table_x
        for c_w in cols:
            p2.draw_line(fitz.Point(curr_x, table_y), fitz.Point(curr_x, table_y + tot_h))
            curr_x += c_w
        p2.draw_line(fitz.Point(table_x + tot_w, table_y), fitz.Point(table_x + tot_w, table_y + tot_h))

        # Insert cell text
        for r_idx, r_vals in enumerate(rows_data):
            cell_y = table_y + r_idx * row_h + 16
            cx = table_x + 6
            for c_idx, val in enumerate(r_vals):
                p2.insert_text((cx, cell_y), val, fontsize=10)
                cx += cols[c_idx]

        doc.save(str(cls.sample_pdf))
        doc.close()

        # 2. Create sample images
        cls.sample_jpg = cls.temp_dir / "sample_img1.jpg"
        img1 = Image.new("RGB", (640, 480), color=(60, 120, 180))
        img1.save(cls.sample_jpg, format="JPEG", quality=90)

        cls.sample_png = cls.temp_dir / "sample_img2.png"
        img2 = Image.new("RGBA", (500, 500), color=(40, 160, 80, 255))
        img2.save(cls.sample_png, format="PNG")

        cls.sample_webp = cls.temp_dir / "sample_img3.webp"
        img3 = Image.new("RGB", (400, 300), color=(200, 90, 40))
        img3.save(cls.sample_webp, format="WEBP", quality=85)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.temp_dir, ignore_errors=True)

    def test_01_conversion_capabilities_matrix(self):
        """Verifies dynamic capability detection and API endpoint."""
        caps = converters.get_conversion_capabilities()
        self.assertIn("engines", caps)
        self.assertIn("conversions", caps)

        # Check essential engines
        self.assertIn("pymupdf", caps["engines"])
        self.assertIn("pillow", caps["engines"])
        self.assertIn("pdf2docx", caps["engines"])
        self.assertIn("openpyxl", caps["engines"])
        self.assertIn("python_docx", caps["engines"])
        self.assertIn("python_pptx", caps["engines"])
        self.assertIn("reportlab", caps["engines"])

        # Check conversion statuses
        self.assertEqual(caps["conversions"]["jpg_to_pdf"]["status"], "FULL")
        self.assertEqual(caps["conversions"]["pdf_to_docx"]["status"], "FULL")
        self.assertEqual(caps["conversions"]["pdf_to_excel"]["status"], "FULL")
        self.assertEqual(caps["conversions"]["pdf_to_powerpoint"]["status"], "FULL")
        self.assertEqual(caps["conversions"]["markdown_to_pdf"]["status"], "FULL")
        self.assertEqual(caps["conversions"]["txt_to_pdf"]["status"], "FULL")

        # Test API endpoint
        res = self.client.get("/api/studio/conversions/capabilities")
        self.assertEqual(res.status_code, 200)
        api_caps = res.json()
        self.assertEqual(api_caps["conversions"]["pdf_to_docx"]["status"], "FULL")

    def test_02_images_to_pdf_single_and_formats(self):
        """Tests image to PDF conversion across JPG, PNG, and WebP with page sizing."""
        out_jpg_pdf = self.temp_dir / "from_jpg.pdf"
        converters.images_to_pdf([self.sample_jpg], out_jpg_pdf, page_size="a4", orientation="portrait", margin_pt=18.0)
        self.assertTrue(out_jpg_pdf.exists())
        doc = fitz.open(str(out_jpg_pdf))
        self.assertEqual(len(doc), 1)
        # Check A4 dimension (595.28 x 841.89)
        self.assertAlmostEqual(doc[0].rect.width, 595.28, delta=1.0)
        doc.close()

        out_png_pdf = self.temp_dir / "from_png.pdf"
        converters.images_to_pdf([self.sample_png], out_png_pdf, page_size="letter", margin_pt=0.0)
        self.assertTrue(out_png_pdf.exists())
        doc_png = fitz.open(str(out_png_pdf))
        self.assertEqual(len(doc_png), 1)
        self.assertAlmostEqual(doc_png[0].rect.width, 612.0, delta=1.0)
        doc_png.close()

        out_webp_pdf = self.temp_dir / "from_webp.pdf"
        converters.images_to_pdf([self.sample_webp], out_webp_pdf, page_size="original")
        self.assertTrue(out_webp_pdf.exists())
        doc_webp = fitz.open(str(out_webp_pdf))
        self.assertEqual(len(doc_webp), 1)
        doc_webp.close()

    def test_03_images_to_pdf_multi_page_and_layout(self):
        """Tests batch multi-image to PDF with ordering, auto orientation, and fit modes."""
        out_multi = self.temp_dir / "multi_combined.pdf"
        images = [self.sample_jpg, self.sample_png, self.sample_webp]
        converters.images_to_pdf(
            images, out_multi,
            page_size="a4",
            orientation="auto",
            margin_pt=24.0,
            fit_mode="fit"
        )
        self.assertTrue(out_multi.exists())
        doc = fitz.open(str(out_multi))
        self.assertEqual(len(doc), 3)
        doc.close()

    def test_04_api_studio_images_to_pdf_endpoint(self):
        """Tests POST /api/studio/convert/images-to-pdf endpoint."""
        files = [
            ("files", ("img1.jpg", self.sample_jpg.read_bytes(), "image/jpeg")),
            ("files", ("img2.png", self.sample_png.read_bytes(), "image/png"))
        ]
        res = self.client.post(
            "/api/studio/convert/images-to-pdf",
            files=files,
            data={"page_size": "a4", "margin": "18.0", "fit_mode": "fit"}
        )
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers.get("content-type"), "application/pdf")
        self.assertEqual(res.headers.get("x-converted-pages"), "2")
        out_doc = fitz.open(stream=res.content, filetype="pdf")
        self.assertEqual(len(out_doc), 2)
        out_doc.close()

    def test_05_pdf_to_images_single_and_zip(self):
        """Tests PDF to JPG: single page direct image and multi-page ZIP bundle."""
        # 1. Single page -> single JPG
        single_out = self.temp_dir / "page1_export.jpg"
        converters.pdf_to_images(self.sample_pdf, single_out, fmt="jpg", dpi=100, pages="1")
        self.assertTrue(single_out.exists())
        with Image.open(single_out) as img:
            self.assertEqual(img.format, "JPEG")
            self.assertGreater(img.width, 0)

        # 2. Multi page -> ZIP archive
        zip_out = self.temp_dir / "all_pages.zip"
        converters.pdf_to_images(self.sample_pdf, zip_out, fmt="jpg", dpi=100, pages="all")
        self.assertTrue(zip_out.exists())
        with zipfile.ZipFile(zip_out, "r") as zf:
            namelist = zf.namelist()
            self.assertEqual(len(namelist), 2)
            self.assertIn("page_001.jpg", namelist)
            self.assertIn("page_002.jpg", namelist)

    def test_06_pdf_to_png_and_webp(self):
        """Tests PDF page export to lossless PNG and WebP formats."""
        png_out = self.temp_dir / "page1.png"
        converters.pdf_to_images(self.sample_pdf, png_out, fmt="png", dpi=120, pages="1")
        self.assertTrue(png_out.exists())
        with Image.open(png_out) as img:
            self.assertEqual(img.format, "PNG")

        webp_out = self.temp_dir / "page1.webp"
        converters.pdf_to_images(self.sample_pdf, webp_out, fmt="webp", dpi=120, pages="1")
        self.assertTrue(webp_out.exists())
        with Image.open(webp_out) as img:
            self.assertEqual(img.format, "WEBP")

    def test_07_text_to_pdf_typography_and_wrapping(self):
        """Tests plain text (.txt) to formatted, paginated PDF."""
        sample_txt = self.temp_dir / "sample_doc.txt"
        sample_txt.write_text(
            "Paperplane Text Document\n\n"
            "This is a paragraph with several lines demonstrating clean typography and automatic wrapping.\n"
            "Paperplane ensures that text converts without truncation or font encoding issues.\n\n"
            "Second Section\n\n"
            "Additional details and bullet points:\n- Point 1\n- Point 2\n- Point 3\n",
            encoding="utf-8"
        )
        out_pdf = self.temp_dir / "from_txt.pdf"
        converters.text_to_pdf(sample_txt, out_pdf, font_size=11, margin_pt=36)
        self.assertTrue(out_pdf.exists())
        doc = fitz.open(str(out_pdf))
        self.assertGreaterEqual(len(doc), 1)
        full_text = doc[0].get_text()
        self.assertIn("Paperplane Text Document", full_text)
        self.assertIn("typography", full_text)
        doc.close()

    def test_08_pdf_to_text_reading_order(self):
        """Tests extracting structured Unicode text from PDF."""
        out_txt = self.temp_dir / "extracted_text.txt"
        converters.pdf_to_text(self.sample_pdf, out_txt, pages="all")
        self.assertTrue(out_txt.exists())
        content = out_txt.read_text(encoding="utf-8")
        self.assertIn("--- Page 1 ---", content)
        self.assertIn("Paperplane Executive Summary 2026", content)
        self.assertIn("--- Page 2 ---", content)
        self.assertIn("Financial Performance", content)

    def test_09_markdown_to_pdf_gfm_tables_and_headings(self):
        """Tests converting Markdown (.md) with GFM tables and code blocks to PDF."""
        sample_md = self.temp_dir / "notes.md"
        sample_md.write_text(
            "# System Architecture\n\n"
            "Here is an overview of the document conversion subsystem.\n\n"
            "| Component | Technology | Status |\n"
            "| --- | --- | --- |\n"
            "| Parser | PyMuPDF / fitz | Production |\n"
            "| Word Engine | pdf2docx | Active |\n"
            "| Spreadsheet | openpyxl | Verified |\n\n"
            "```python\ndef convert(src, dst):\n    pass\n```\n\n"
            "> Built for enterprise-grade performance.\n",
            encoding="utf-8"
        )
        out_pdf = self.temp_dir / "from_md.pdf"
        converters.markdown_to_pdf(sample_md, out_pdf)
        self.assertTrue(out_pdf.exists())
        doc = fitz.open(str(out_pdf))
        self.assertGreaterEqual(len(doc), 1)
        text = doc[0].get_text()
        self.assertIn("System Architecture", text)
        self.assertIn("PyMuPDF", text)
        doc.close()

    def test_10_pdf_to_markdown_structured_extraction(self):
        """Tests extracting semantic Markdown preserving headings, lists, and tables."""
        out_md = self.temp_dir / "extracted_doc.md"
        converters.pdf_to_markdown(self.sample_pdf, out_md, pages="all")
        self.assertTrue(out_md.exists())
        content = out_md.read_text(encoding="utf-8")
        self.assertIn("# Paperplane Executive Summary 2026", content)
        self.assertIn("Financial Performance", content)
        self.assertIn("| Metric | Target | Achieved |", content)

    def test_11_html_to_pdf_rendering(self):
        """Tests converting HTML snippet into PDF."""
        html_code = """
        <!DOCTYPE html>
        <html>
        <head><title>Invoice #1042</title></head>
        <body>
        <h1>Invoice #1042</h1>
        <p>Customer: ACME Corporation</p>
        <table border="1">
            <tr><th>Item</th><th>Total</th></tr>
            <tr><td>Paperplane License</td><td>$99.00</td></tr>
        </table>
        </body>
        </html>
        """
        out_pdf = self.temp_dir / "from_html.pdf"
        converters.html_to_pdf(html_code, out_pdf, is_url=False)
        self.assertTrue(out_pdf.exists())
        doc = fitz.open(str(out_pdf))
        self.assertGreaterEqual(len(doc), 1)
        text = doc[0].get_text()
        self.assertIn("Invoice #1042", text)
        self.assertIn("ACME Corporation", text)
        doc.close()

    def test_12_html_ssrf_security_guards(self):
        """Verifies SSRF security validator blocks private, loopback, and metadata IPs."""
        # 1. Loopback hostname
        safe, reason = converters.is_safe_url("http://localhost:8501/api/secret")
        self.assertFalse(safe)
        self.assertIn("localhost", reason.lower())

        # 2. Local IP 127.0.0.1
        safe, reason = converters.is_safe_url("http://127.0.0.1:8000/")
        self.assertFalse(safe)

        # 3. AWS / Cloud metadata endpoint
        safe, reason = converters.is_safe_url("http://169.254.169.254/latest/meta-data/")
        self.assertFalse(safe)

        # 4. Dangerous scheme
        safe, reason = converters.is_safe_url("file:///etc/passwd")
        self.assertFalse(safe)

        # 5. Verify html_to_pdf raises ValueError on unsafe URL
        out_pdf = self.temp_dir / "blocked_ssrf.pdf"
        with self.assertRaises(ValueError) as ctx:
            converters.html_to_pdf("http://127.0.0.1:8501/admin", out_pdf, is_url=True)
        self.assertIn("Security Guard", str(ctx.exception))

    def test_13_pdf_to_html_sanitized_document(self):
        """Tests converting PDF to standalone, sanitized HTML with page sections."""
        out_html = self.temp_dir / "converted_page.html"
        converters.pdf_to_html(self.sample_pdf, out_html, pages="all")
        self.assertTrue(out_html.exists())
        content = out_html.read_text(encoding="utf-8")
        self.assertIn('<section class="pdf-page" id="page-1">', content)
        self.assertIn('<h1>Paperplane Executive Summary 2026</h1>', content)
        self.assertIn('<section class="pdf-page" id="page-2">', content)
        self.assertIn('<table>', content)
        self.assertIn('<th>Metric</th>', content)
        # Check no unsanitized script tags
        self.assertNotIn('<script>', content.lower())

    def test_14_word_to_pdf_and_pdf_to_word(self):
        """Tests bidirectional Word (.docx) <-> PDF conversions."""
        # Create Word document
        sample_docx = self.temp_dir / "contract.docx"
        doc = docx.Document()
        doc.add_heading("Service Agreement", level=1)
        doc.add_paragraph("This agreement is made between Party A and Party B.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Term"
        table.cell(0, 1).text = "Details"
        table.cell(1, 0).text = "Duration"
        table.cell(1, 1).text = "12 Months"
        doc.save(str(sample_docx))

        # DOCX -> PDF
        out_pdf = self.temp_dir / "contract.pdf"
        converters.word_to_pdf(sample_docx, out_pdf)
        self.assertTrue(out_pdf.exists())
        pdf_doc = fitz.open(str(out_pdf))
        self.assertGreaterEqual(len(pdf_doc), 1)
        self.assertIn("Service Agreement", pdf_doc[0].get_text())
        pdf_doc.close()

        # PDF -> DOCX
        back_docx = self.temp_dir / "contract_back.docx"
        converters.pdf_to_word(out_pdf, back_docx)
        self.assertTrue(back_docx.exists())
        d = docx.Document(str(back_docx))
        p_texts = [p.text for p in d.paragraphs if p.text]
        self.assertTrue(any("Service Agreement" in t for t in p_texts))

    def test_15_excel_to_pdf_and_pdf_to_excel(self):
        """Tests bidirectional Excel (.xlsx) <-> PDF conversions with tables."""
        # Create Excel document
        sample_xlsx = self.temp_dir / "inventory.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Stock"
        ws.append(["SKU", "Item Description", "Quantity", "Unit Cost"])
        ws.append(["A001", "Precision Screwdriver", 150, 4.50])
        ws.append(["A002", "Hex Nut Pack 100pc", 80, 2.20])
        wb.save(str(sample_xlsx))

        # XLSX -> PDF
        out_pdf = self.temp_dir / "inventory.pdf"
        converters.excel_to_pdf(sample_xlsx, out_pdf)
        self.assertTrue(out_pdf.exists())
        pdf_doc = fitz.open(str(out_pdf))
        self.assertGreaterEqual(len(pdf_doc), 1)
        pdf_doc.close()

        # PDF -> XLSX
        back_xlsx = self.temp_dir / "inventory_back.xlsx"
        converters.pdf_to_excel(self.sample_pdf, back_xlsx)
        self.assertTrue(back_xlsx.exists())
        wb_back = openpyxl.load_workbook(str(back_xlsx), data_only=True)
        self.assertGreaterEqual(len(wb_back.sheetnames), 1)
        wb_back.close()

    def test_16_powerpoint_to_pdf_and_pdf_to_powerpoint(self):
        """Tests bidirectional PowerPoint (.pptx) <-> PDF conversions in visual and editable modes."""
        sample_pptx = self.temp_dir / "pitch.pptx"
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Paperplane Series A"
        slide.placeholders[1].text = "Revolutionizing offline document workflows."
        prs.save(str(sample_pptx))

        # PPTX -> PDF
        out_pdf = self.temp_dir / "pitch.pdf"
        converters.powerpoint_to_pdf(sample_pptx, out_pdf)
        self.assertTrue(out_pdf.exists())
        doc = fitz.open(str(out_pdf))
        self.assertEqual(len(doc), 1)
        self.assertIn("Paperplane Series A", doc[0].get_text())
        doc.close()

        # PDF -> PPTX (Visual mode)
        back_pptx = self.temp_dir / "pitch_back_visual.pptx"
        converters.pdf_to_powerpoint(self.sample_pdf, back_pptx, mode="visual")
        self.assertTrue(back_pptx.exists())
        p_vis = pptx.Presentation(str(back_pptx))
        self.assertEqual(len(p_vis.slides), 2)

        # PDF -> PPTX (Editable mode)
        back_editable = self.temp_dir / "pitch_back_editable.pptx"
        converters.pdf_to_powerpoint(self.sample_pdf, back_editable, mode="editable")
        self.assertTrue(back_editable.exists())
        p_edit = pptx.Presentation(str(back_editable))
        self.assertEqual(len(p_edit.slides), 2)
        # Check shapes on slide 1 include text boxes
        shapes = p_edit.slides[0].shapes
        self.assertGreater(len(shapes), 1)

    def test_17_pdf_to_pdfa_archival(self):
        """Tests PDF/A-1b metadata and standardization."""
        out_pdfa = self.temp_dir / "archived.pdf"
        converters.pdf_to_pdfa(self.sample_pdf, out_pdfa)
        self.assertTrue(out_pdfa.exists())
        doc = fitz.open(str(out_pdfa))
        self.assertEqual(len(doc), 2)
        meta = doc.metadata or {}
        self.assertEqual(meta.get("keywords"), "PDF/A-1b")
        self.assertEqual(meta.get("producer"), "Paperplane PDF/A Archival Engine")
        doc.close()

    def test_18_validate_conversion_output_suite(self):
        """Tests validation engine on valid formats and safe error detection."""
        # Valid PDF
        rep = converters.validate_conversion_output(self.sample_pdf, "pdf")
        self.assertTrue(rep["valid"])
        self.assertEqual(rep["metrics"]["pages"], 2)

        # Nonexistent file
        with self.assertRaises(ValueError):
            converters.validate_conversion_output(self.temp_dir / "ghost.pdf", "pdf")

        # Corrupted / invalid PDF signature
        corrupt = self.temp_dir / "corrupt.pdf"
        corrupt.write_bytes(b"NOT_A_PDF_FILE")
        with self.assertRaises(ValueError) as ctx:
            converters.validate_conversion_output(corrupt, "pdf")
        self.assertIn("magic bytes", str(ctx.exception).lower())

    def test_19_api_studio_convert_endpoint(self):
        """Tests unified studio conversion endpoint across multiple target formats."""
        # 1. Convert PDF -> Word (.docx)
        res_docx = self.client.post(
            "/api/studio/convert",
            files={"file": ("sample.pdf", self.sample_pdf.read_bytes(), "application/pdf")},
            data={"target_format": "docx", "options": json.dumps({"pages": "all"})}
        )
        self.assertEqual(res_docx.status_code, 200)
        self.assertIn("wordprocessingml", res_docx.headers.get("content-type", ""))
        self.assertEqual(res_docx.headers.get("x-conversion-valid"), "True")

        # 2. Convert PDF -> Markdown (.md)
        res_md = self.client.post(
            "/api/studio/convert",
            files={"file": ("sample.pdf", self.sample_pdf.read_bytes(), "application/pdf")},
            data={"target_format": "md"}
        )
        self.assertEqual(res_md.status_code, 200)
        self.assertIn("markdown", res_md.headers.get("content-type", ""))
        self.assertIn("Paperplane Executive Summary", res_md.content.decode("utf-8"))

        # 3. Convert PDF -> Text (.txt)
        res_txt = self.client.post(
            "/api/studio/convert",
            files={"file": ("sample.pdf", self.sample_pdf.read_bytes(), "application/pdf")},
            data={"target_format": "txt"}
        )
        self.assertEqual(res_txt.status_code, 200)
        self.assertIn("text/plain", res_txt.headers.get("content-type", ""))

    def test_20_process_job_queue_conversions(self):
        """Verifies asynchronous job queue process() execution for Phase 5 operations."""
        job_dir = self.temp_dir / "job_test"
        job_dir.mkdir(parents=True, exist_ok=True)

        # 1. Process images to PDF via process()
        out1 = server.process("jpg_to_pdf", [self.sample_jpg, self.sample_png], {"page_size": "a4"}, job_dir)
        self.assertTrue(out1.exists())
        self.assertEqual(len(fitz.open(str(out1))), 2)

        # 2. Process text to PDF
        sample_txt = self.temp_dir / "job_note.txt"
        sample_txt.write_text("Hello Queue Worker", encoding="utf-8")
        out2 = server.process("text_to_pdf", [sample_txt], {}, job_dir)
        self.assertTrue(out2.exists())
        self.assertEqual(len(fitz.open(str(out2))), 1)

        # 3. Process PDF to Markdown
        out3 = server.process("pdf_to_markdown", [self.sample_pdf], {}, job_dir)
        self.assertTrue(out3.exists())
        self.assertTrue(out3.name.endswith(".md"))


if __name__ == "__main__":
    unittest.main()
